from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

from app.parse_docs import _docx_paragraph_to_markdown, extract_text_from_file


def add_hyperlink(paragraph, text: str) -> None:
    """Build a `<w:hyperlink><w:r>` wrapper, which python-docx cannot author."""
    hyperlink = paragraph._p.makeelement(qn("w:hyperlink"), {})
    run = paragraph._p.makeelement(qn("w:r"), {})
    text_node = paragraph._p.makeelement(qn("w:t"), {})
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _png_bytes(*, width: int = 320, height: int = 240, color=(40, 120, 200)) -> bytes:
    image = Image.new("RGB", (width, height), color=color)
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def test_docx_runs_become_inline_markdown():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Plain ")
    bold = paragraph.add_run("bold")
    bold.bold = True
    paragraph.add_run(", ")
    italic = paragraph.add_run("italic")
    italic.italic = True
    paragraph.add_run(", and ")
    both = paragraph.add_run("both")
    both.bold = True
    both.italic = True
    paragraph.add_run(".")

    assert _docx_paragraph_to_markdown(paragraph) == (
        "Plain **bold**, *italic*, and ***both***."
    )


def test_adjacent_runs_with_same_style_are_merged():
    document = Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("two ")
    first.bold = True
    second = paragraph.add_run("runs")
    second.bold = True

    assert _docx_paragraph_to_markdown(paragraph) == "**two runs**"


def test_literal_markdown_characters_are_escaped():
    document = Document()
    paragraph = document.add_paragraph(r"Use * literally and keep \ paths")

    assert _docx_paragraph_to_markdown(paragraph) == (
        r"Use \* literally and keep \\ paths"
    )


def test_hyperlink_text_is_kept():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("See ")
    add_hyperlink(paragraph, "this page")
    paragraph.add_run(" for details.")

    assert _docx_paragraph_to_markdown(paragraph) == "See this page for details."


def test_paragraph_made_only_of_a_hyperlink_is_not_dropped():
    document = Document()
    paragraph = document.add_paragraph()
    add_hyperlink(paragraph, "Entirely a hyperlink.")

    assert _docx_paragraph_to_markdown(paragraph) == "Entirely a hyperlink."


def test_extract_includes_table_cells(tmp_path):
    document = Document()
    document.add_paragraph("Body paragraph.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "First cell."
    table.rows[0].cells[1].text = "Second cell."
    path = tmp_path / "with-table.docx"
    document.save(path)

    text = extract_text_from_file(path, "with-table.docx")

    assert text.split("\n\n") == ["Body paragraph.", "First cell.", "Second cell."]


def test_extract_keeps_images_and_caption_style(tmp_path):
    png_path = tmp_path / "figure.png"
    png_path.write_bytes(_png_bytes())

    document = Document()
    document.add_paragraph("Before the figure.")
    picture = document.add_paragraph()
    picture.add_run().add_picture(str(png_path), width=Inches(3))
    caption = document.add_paragraph("Figure 1. The harbor at dawn.")
    try:
        caption.style = "Caption"
    except KeyError:
        # Minimal test docs may lack the built-in Caption style; set pStyle directly.
        p_pr = caption._p.get_or_add_pPr()
        style_el = p_pr.makeelement(qn("w:pStyle"), {qn("w:val"): "Caption"})
        p_pr.append(style_el)
    document.add_paragraph("After the figure.")

    path = tmp_path / "with-figure.docx"
    document.save(path)
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()

    text = extract_text_from_file(
        path,
        "with-figure.docx",
        media_dir=upload_dir,
        book_id="bookfig1",
    )
    blocks = text.split("\n\n")

    assert blocks[0] == "Before the figure."
    assert blocks[1].startswith("![Figure 1. The harbor at dawn.](/api/books/bookfig1/media/")
    assert blocks[1].endswith(".jpg)")
    assert blocks[2] == "After the figure."

    asset_name = blocks[1].rsplit("/", 1)[-1].rstrip(")")
    assert (upload_dir / "media" / "bookfig1" / asset_name).is_file()


def test_extract_without_media_dir_skips_images(tmp_path):
    png_path = tmp_path / "figure.png"
    png_path.write_bytes(_png_bytes())

    document = Document()
    document.add_paragraph("Only text survives.")
    document.add_paragraph().add_run().add_picture(str(png_path), width=Inches(2))
    path = tmp_path / "img.docx"
    document.save(path)

    text = extract_text_from_file(path, "img.docx")

    assert text == "Only text survives."
    assert "![" not in text
