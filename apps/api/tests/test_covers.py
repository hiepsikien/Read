import io
from datetime import datetime, timezone
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from nanoid import generate
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.auth import hash_password, mint_dev_id_token
from app.categories import ensure_categories
from app.config import Settings
from app.covers import (
    extract_cover_from_docx,
    normalize_cover_image,
    save_cover_bytes,
    try_extract_and_save_cover,
)
from app.db import Base, get_db
from app.main import app
from app.models import Book, User
from app.routers import books as books_router


@pytest.fixture()
def db_session():
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(bind=engine)
    session = Session()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture()
def client(db_session):
    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


@pytest.fixture()
def seeded(db_session):
    categories = ensure_categories(db_session)
    now = datetime.now(timezone.utc)
    publisher = User(
        id=generate(),
        firebase_uid=f"dev-{generate()}",
        email="author@example.com",
        name="Author",
        role="publisher",
        password_hash=hash_password("password123"),
        created_at=now,
    )
    reader = User(
        id=generate(),
        firebase_uid=f"dev-{generate()}",
        email="reader@example.com",
        name="Reader",
        role="reader",
        password_hash=hash_password("reader123"),
        created_at=now,
    )
    db_session.add_all([publisher, reader])
    db_session.commit()
    return {
        "categories": categories,
        "publisher": publisher,
        "reader": reader,
        "fiction": next(c for c in categories if c.slug == "fiction"),
    }


def auth_header(user: User) -> dict[str, str]:
    token = mint_dev_id_token(
        uid=user.firebase_uid or f"dev-{user.id}",
        email=user.email,
        name=user.name,
    )
    return {"Authorization": f"Bearer {token}"}


def _make_image_bytes(*, width: int, height: int, color=(40, 90, 70), fmt="JPEG") -> bytes:
    image = Image.new("RGB", (width, height), color)
    buffer = io.BytesIO()
    image.save(buffer, format=fmt, quality=90)
    return buffer.getvalue()


def _large_cover_bytes(*, width: int = 800, height: int = 1200) -> bytes:
    image = Image.effect_noise((width, height), 48).convert("RGB")
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=95)
    data = buffer.getvalue()
    assert len(data) >= 80 * 1024
    return data


def _docx_with_image(path: Path, image_bytes: bytes, filename: str = "cover.jpg") -> None:
    document = Document()
    document.add_paragraph("Title page")
    image_path = path.parent / filename
    image_path.write_bytes(image_bytes)
    document.add_picture(str(image_path))
    document.add_paragraph("Chapter one body text for extraction.")
    document.save(str(path))


def test_normalize_cover_image_reencodes_and_resizes():
    raw = _make_image_bytes(width=2400, height=3600)
    out = normalize_cover_image(raw)
    image = Image.open(io.BytesIO(out))
    assert image.format == "JPEG"
    assert max(image.size) == 1600


def test_extract_cover_from_docx_picks_qualifying_image(tmp_path: Path):
    docx_path = tmp_path / "with-cover.docx"
    cover = _large_cover_bytes()
    _docx_with_image(docx_path, cover)
    extracted = extract_cover_from_docx(docx_path)
    assert extracted is not None
    assert len(extracted) >= 80 * 1024


def test_extract_cover_skips_tiny_images(tmp_path: Path):
    docx_path = tmp_path / "tiny.docx"
    tiny = _make_image_bytes(width=64, height=64)
    _docx_with_image(docx_path, tiny, filename="icon.jpg")
    assert extract_cover_from_docx(docx_path) is None


def test_try_extract_fail_soft(tmp_path: Path):
    missing = tmp_path / "missing.docx"
    assert try_extract_and_save_cover(tmp_path, "book1", missing) is None


def test_create_book_extracts_cover(client, seeded, tmp_path, monkeypatch):
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    settings = Settings(upload_dir=str(upload_dir), auth_dev_mode=True)
    monkeypatch.setattr(books_router, "get_settings", lambda: settings)

    docx_path = tmp_path / "ms.docx"
    _docx_with_image(docx_path, _large_cover_bytes())

    headers = auth_header(seeded["publisher"])
    with docx_path.open("rb") as handle:
        response = client.post(
            "/api/books",
            data={
                "title": "Covered Book",
                "description": "Has a cover",
                "pricing": "free",
                "category_id": seeded["fiction"].id,
            },
            files={
                "file": (
                    "ms.docx",
                    handle,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=headers,
        )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["id"]
    assert payload["cover_url"] == f"/api/books/{payload['id']}/cover"

    cover_resp = client.get(payload["cover_url"], headers=headers)
    assert cover_resp.status_code == 200
    assert cover_resp.headers["content-type"].startswith("image/")


def test_upload_cover_override_and_access(client, seeded, tmp_path, monkeypatch, db_session):
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    settings = Settings(upload_dir=str(upload_dir), auth_dev_mode=True)
    monkeypatch.setattr(books_router, "get_settings", lambda: settings)

    now = datetime.now(timezone.utc)
    book = Book(
        id=generate(),
        publisher_id=seeded["publisher"].id,
        category_id=seeded["fiction"].id,
        title="Draft Cover",
        description="",
        price_cents=0,
        status="draft",
        cover_path=None,
        created_at=now,
        updated_at=now,
    )
    db_session.add(book)
    db_session.commit()

    assert client.get(f"/api/books/{book.id}/cover").status_code == 404

    png = _make_image_bytes(width=700, height=1000, fmt="PNG")
    headers = auth_header(seeded["publisher"])
    upload = client.post(
        f"/api/books/{book.id}/cover",
        files={"file": ("cover.png", png, "image/png")},
        headers=headers,
    )
    assert upload.status_code == 200, upload.text
    assert upload.json()["cover_url"] == f"/api/books/{book.id}/cover"

    assert client.get(f"/api/books/{book.id}/cover", headers=headers).status_code == 200
    stranger_headers = auth_header(seeded["reader"])
    assert client.get(f"/api/books/{book.id}/cover", headers=stranger_headers).status_code == 404

    book.status = "published"
    db_session.commit()
    assert client.get(f"/api/books/{book.id}/cover").status_code == 200

    book.status = "pending_review"
    db_session.commit()
    blocked = client.post(
        f"/api/books/{book.id}/cover",
        files={"file": ("cover.png", png, "image/png")},
        headers=headers,
    )
    assert blocked.status_code == 400


def test_list_includes_cover_url(client, seeded, db_session, tmp_path, monkeypatch):
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    settings = Settings(upload_dir=str(upload_dir), auth_dev_mode=True)
    monkeypatch.setattr(books_router, "get_settings", lambda: settings)

    book_id = generate()
    rel = save_cover_bytes(upload_dir, book_id, _make_image_bytes(width=500, height=750))
    now = datetime.now(timezone.utc)
    book = Book(
        id=book_id,
        publisher_id=seeded["publisher"].id,
        category_id=seeded["fiction"].id,
        title="Listed",
        description="d",
        price_cents=0,
        status="published",
        cover_path=rel,
        created_at=now,
        updated_at=now,
    )
    db_session.add(book)
    db_session.commit()

    response = client.get("/api/books")
    assert response.status_code == 200
    found = next(b for b in response.json()["books"] if b["id"] == book_id)
    assert found["cover_url"] == f"/api/books/{book_id}/cover"
