from datetime import datetime, timezone

import pytest
from fastapi.testclient import TestClient
from nanoid import generate
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.auth import hash_password, mint_dev_id_token
from app.categories import ensure_categories
from app.db import Base, get_db
from app.glossary import aliases_to_storage
from app.main import app
from app.models import Book, Chapter, GlossaryEntry, User


@pytest.fixture()
def db_session():
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(bind=engine)
    session = Session()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture()
def client(db_session):
    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.router.on_startup.clear()
    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


@pytest.fixture()
def seeded(db_session):
    categories = ensure_categories(db_session)
    now = datetime.now(timezone.utc)
    publisher = User(
        id=generate(),
        firebase_uid=f"dev-{generate()}",
        email="author@example.com",
        name="Author",
        role="publisher",
        password_hash=hash_password("password123"),
        created_at=now,
    )
    db_session.add(publisher)
    book = Book(
        id=generate(),
        publisher_id=publisher.id,
        category_id=categories[0].id,
        title="S1E1 — Đại Lộ Đại Dương",
        description="Sample",
        price_cents=0,
        status="published",
        raw_text="ok",
        created_at=now,
        updated_at=now,
    )
    db_session.add(book)
    chapter = Chapter(
        id=generate(),
        book_id=book.id,
        position=1,
        title="S1E1C1 — Opening",
        content=(
            "Bờ biển Calicut.\n\n"
            "VASCO DA GAMA đã đi vào lịch sử.\n\n"
            "AFONSO DE ALBUQUERQUE đứng trên boong."
        ),
        word_count=20,
        group_index=1,
    )
    db_session.add(chapter)
    for name, aliases, summary in [
        ("Vasco da Gama", [], "(k. 1469 – 1524) Nhà thám hiểm Bồ Đào Nha."),
        ("Afonso de Albuquerque", [], "(k. 1453 – 1515) Đô đốc Bồ Đào Nha."),
        ("Mạc Đăng Dung", [], "(1483 – 1541) Người làng Cổ Trai."),
    ]:
        db_session.add(
            GlossaryEntry(
                id=generate(),
                book_id=book.id,
                episode_key="S1E1",
                episode_title="GIÔNG BÃO KINH THÀNH",
                group_label="NHÂN VẬT LỊCH SỬ QUỐC TẾ",
                name=name,
                aliases=aliases_to_storage(aliases),
                summary=summary,
                sort_key=name.casefold(),
                created_at=now,
                updated_at=now,
            )
        )
    db_session.commit()
    return {"publisher": publisher, "book": book, "chapter": chapter, "db": db_session}


def test_explain_paid_later_chapter_does_not_require_purchase(client, seeded):
    from app.glossary import aliases_to_storage
    from app.models import Chapter, GlossaryEntry

    book = seeded["book"]
    book.price_cents = 499
    now = datetime.now(timezone.utc)
    paid_chapter = Chapter(
        id=generate(),
        book_id=book.id,
        position=2,
        title="Chapter II",
        content="Seneca[9] on the later argument.",
        word_count=6,
        group_index=2,
    )
    note = GlossaryEntry(
        id=generate(),
        book_id=book.id,
        episode_key="II",
        episode_title="Seneca",
        group_label="Chú thích",
        name="Seneca [9]",
        aliases=aliases_to_storage(["[9]"]),
        summary="Seneca, Natural Questions.",
        sort_key="seneca [9]",
        created_at=now,
        updated_at=now,
    )
    seeded["db"].add_all([paid_chapter, note])
    seeded["db"].commit()

    response = client.post(
        f"/api/books/{book.id}/chapters/{paid_chapter.id}/explain",
        json={"entry_id": note.id},
    )
    assert response.status_code == 200, response.text
    assert response.json()["card"]["title"] == "Seneca [9]"


def test_explain_returns_book_note_without_ai(client, seeded):
    book = seeded["book"]
    chapter = seeded["chapter"]
    response = client.post(
        f"/api/books/{book.id}/chapters/{chapter.id}/explain",
        json={"query": "Vasco da Gama", "need_context": False},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "ok"
    assert payload["ai_used"] is False
    assert payload["card"]["title"] == "Vasco da Gama"
    assert "Bồ Đào Nha" in payload["card"]["book_note"]
    assert "book" in payload["card"]["sources"]
    assert payload["card"]["ai_context"] == ""


def test_explain_cache_hit(client, seeded):
    book = seeded["book"]
    chapter = seeded["chapter"]
    path = f"/api/books/{book.id}/chapters/{chapter.id}/explain"
    first = client.post(path, json={"query": "Vasco da Gama"}).json()
    second = client.post(path, json={"query": "Vasco da Gama"}).json()
    assert first["cache_hit"] is False
    assert second["cache_hit"] is True


def test_explain_note_cached_across_chapters(client, seeded):
    from app.glossary import aliases_to_storage
    from app.models import Chapter, GlossaryEntry

    book = seeded["book"]
    chapter = seeded["chapter"]
    now = datetime.now(timezone.utc)
    note = GlossaryEntry(
        id=generate(),
        book_id=book.id,
        episode_key="I",
        episode_title="Seneca",
        group_label="Chú thích",
        name="Seneca [9]",
        aliases=aliases_to_storage(["[9]"]),
        summary="Seneca, Natural Questions.",
        sort_key="seneca [9]",
        created_at=now,
        updated_at=now,
    )
    other = Chapter(
        id=generate(),
        book_id=book.id,
        position=2,
        title="Chapter II",
        content="Another page.",
        word_count=2,
        group_index=1,
    )
    seeded["db"].add_all([note, other])
    seeded["db"].commit()

    first = client.post(
        f"/api/books/{book.id}/chapters/{chapter.id}/explain",
        json={"entry_id": note.id},
    ).json()
    second = client.post(
        f"/api/books/{book.id}/chapters/{other.id}/explain",
        json={"entry_id": note.id},
    ).json()
    assert first["status"] == "ok"
    assert first["cache_hit"] is False
    assert second["cache_hit"] is True
    assert second["card"]["title"] == "Seneca [9]"
    assert "Seneca" in second["card"]["book_note"]


def test_explain_paragraph_notes_follow_marker_order(client, seeded):
    from app.glossary import aliases_to_storage
    from app.models import Chapter, GlossaryEntry

    book = seeded["book"]
    now = datetime.now(timezone.utc)
    chapter = Chapter(
        id=generate(),
        book_id=book.id,
        position=3,
        title="Notes",
        content="Moses[11] then Augustine[12] and Baldus[14].",
        word_count=6,
        group_index=1,
    )
    seeded["db"].add(chapter)
    for name, marker in [("Augustine [12]", "[12]"), ("Baldus [14]", "[14]"), ("Moses [11]", "[11]")]:
        seeded["db"].add(
            GlossaryEntry(
                id=generate(),
                book_id=book.id,
                episode_key="",
                episode_title=name.split(" [")[0],
                group_label="Chú thích",
                name=name,
                aliases=aliases_to_storage([marker]),
                summary=name,
                sort_key=name.casefold(),
                created_at=now,
                updated_at=now,
            )
        )
    seeded["db"].commit()
    payload = client.post(
        f"/api/books/{book.id}/chapters/{chapter.id}/explain",
        json={"paragraph_index": 0},
    ).json()
    assert payload["status"] == "ok"
    assert payload["card"]["glossary_entry"] is None
    assert payload["card"]["ai_context"] == ""
    assert payload["ai_used"] is False
    assert [item["name"] for item in payload["candidates"]] == [
        "Moses [11]",
        "Augustine [12]",
        "Baldus [14]",
    ]
    assert [item["summary"] for item in payload["candidates"]] == [
        "Moses [11]",
        "Augustine [12]",
        "Baldus [14]",
    ]


def test_explain_paragraph_without_notes(client, seeded):
    book = seeded["book"]
    chapter = seeded["chapter"]
    response = client.post(
        f"/api/books/{book.id}/chapters/{chapter.id}/explain",
        json={"paragraph_index": 0},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "ok"
    assert payload["card"]["title"] == "Bờ biển Calicut."
    assert payload["card"]["book_note"] == ""
    assert payload["card"]["ai_context"] == ""
    assert payload["card"]["glossary_entry"] is None
    assert payload["candidates"] == []
    assert payload["ai_used"] is False


def test_explain_paragraph_candidates(client, seeded):
    book = seeded["book"]
    chapter = seeded["chapter"]
    # Paragraph 1 mentions Vasco — keep the note as a candidate, explain the passage.
    response = client.post(
        f"/api/books/{book.id}/chapters/{chapter.id}/explain",
        json={"paragraph_index": 1},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "ok"
    assert payload["card"]["glossary_entry"] is None
    assert payload["card"]["ai_context"] == ""
    assert "VASCO DA GAMA" in payload["card"]["title"]
    assert [item["name"] for item in payload["candidates"]] == ["Vasco da Gama"]
    assert "Bồ Đào Nha" in payload["candidates"][0]["summary"]


def test_list_glossary_compact(client, seeded):
    book = seeded["book"]
    response = client.get(f"/api/books/{book.id}/glossary")
    assert response.status_code == 200
    payload = response.json()
    assert payload["count"] == 3
    assert "summary" not in payload["entries"][0]


def test_upload_glossary_docx(client, seeded, tmp_path):
    from docx import Document

    book = seeded["book"]
    book.status = "draft"
    seeded["db"].commit()
    token = mint_dev_id_token(
        uid=seeded["publisher"].firebase_uid or seeded["publisher"].id,
        email=seeded["publisher"].email,
        name=seeded["publisher"].name,
    )

    doc = Document()
    doc.add_heading("NHÂN VẬT", level=1)
    doc.add_heading("S1E1: TEST", level=2)
    doc.add_paragraph("NHÂN VẬT LỊCH SỬ QUỐC TẾ")
    doc.add_paragraph("Christopher Columbus (1451 – 1506): Nhà hàng hải người Ý.")
    path = tmp_path / "glossary.docx"
    doc.save(path)

    with path.open("rb") as handle:
        response = client.post(
            f"/api/books/{book.id}/glossary",
            headers={"Authorization": f"Bearer {token}"},
            files={
                "file": (
                    "glossary.docx",
                    handle,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["ok"] is True
    assert payload["count"] == 1

    listed = client.get(
        f"/api/books/{book.id}/glossary",
        headers={"Authorization": f"Bearer {token}"},
    ).json()
    assert listed["count"] == 1
    assert listed["entries"][0]["name"] == "Christopher Columbus"


def test_explain_editorial_note_skips_ai_until_asked(client, seeded):
    from app.glossary import aliases_to_storage
    from app.models import GlossaryEntry

    book = seeded["book"]
    chapter = seeded["chapter"]
    now = datetime.now(timezone.utc)
    note = GlossaryEntry(
        id=generate(),
        book_id=book.id,
        episode_key="",
        episode_title="chiếm hữu thực tế",
        group_label="Bối cảnh",
        name="Bối cảnh pháp lý",
        aliases=aliases_to_storage(["chiếm hữu thực tế"]),
        summary="Possessio cần corpus và animus.",
        sort_key="boi canh phap ly",
        created_at=now,
        updated_at=now,
    )
    seeded["db"].add(note)
    seeded["db"].commit()

    first = client.post(
        f"/api/books/{book.id}/chapters/{chapter.id}/explain",
        json={"entry_id": note.id},
    ).json()
    assert first["status"] == "ok"
    assert first["card"]["title"] == "chiếm hữu thực tế"
    assert "Possessio" in first["card"]["book_note"]
    assert first["card"]["ai_context"] == ""
    assert first["ai_used"] is False

    second = client.post(
        f"/api/books/{book.id}/chapters/{chapter.id}/explain",
        json={"entry_id": note.id, "need_context": True},
    ).json()
    assert second["card"]["title"] == "chiếm hữu thực tế"
    assert "Possessio" in second["card"]["book_note"]


def test_paragraph_card_title_truncates_on_word():
    from app.explain import paragraph_card_title

    assert paragraph_card_title("Bờ biển Calicut.") == "Bờ biển Calicut."
    long = " ".join(["word"] * 40)
    title = paragraph_card_title(long)
    assert title.endswith("…")
    assert len(title) <= 73
